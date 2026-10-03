"""
plugins/pptx_plugin.py — Plugin de Generación de Presentaciones (PowerPoint / Google Slides)

100% Local, sin costos ni límites de API. Crea archivos .pptx modernos y estéticos
en el directorio sandbox (W:\\prcts\\jarvis_proyect\\output).
"""
import os
import json
from typing import Dict, Any, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from plugins.base_plugin import BasePlugin
from config import SANDBOX_DIR
from data.sql_store import log_accion


# Paletas cromáticas para diseño limpio
TEMAS_COLOR = {
    "azul_moderno": {
        "fondo": RGBColor(245, 247, 250),
        "titulo": RGBColor(16, 42, 77),
        "texto": RGBColor(51, 65, 85),
        "acento": RGBColor(37, 99, 235),
    },
    "cyber_dark": {
        "fondo": RGBColor(15, 23, 42),
        "titulo": RGBColor(56, 189, 248),
        "texto": RGBColor(226, 232, 240),
        "acento": RGBColor(168, 85, 247),
    },
    "esmeralda": {
        "fondo": RGBColor(240, 253, 244),
        "titulo": RGBColor(6, 78, 59),
        "texto": RGBColor(30, 41, 59),
        "acento": RGBColor(16, 185, 129),
    },
}


class PptxPlugin(BasePlugin):
    @property
    def plugin_id(self) -> str:
        return "pptx_generator"

    @property
    def name(self) -> str:
        return "Generador de Presentaciones (PPTX)"

    @property
    def description(self) -> str:
        return "Crea presentaciones PowerPoint (.pptx) profesionales y editables en el sandbox local."

    def get_tools(self) -> List[Dict[str, Any]]:
        return [
            {
                "type": "function",
                "function": {
                    "name": "crear_presentacion_pptx",
                    "description": "Crea una presentación editable en formato .pptx con portada y diapositivas de contenido formateadas.",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "nombre_archivo": {
                                "type": "string",
                                "description": "Nombre del archivo (ej: 'estrategia_ia.pptx')."
                            },
                            "titulo_principal": {
                                "type": "string",
                                "description": "Título de la presentación en la portada."
                            },
                            "subtitulo": {
                                "type": "string",
                                "description": "Subtítulo o autor de la portada."
                            },
                            "diapositivas_json": {
                                "type": "string",
                                "description": "JSON array con la estructura de diapositivas: [{\"titulo\": \"...\", \"puntos\": [\"punto 1\", \"punto 2\"]}]"
                            },
                            "tema": {
                                "type": "string",
                                "enum": ["azul_moderno", "cyber_dark", "esmeralda"],
                                "default": "azul_moderno",
                                "description": "Estilo visual de la presentación."
                            }
                        },
                        "required": ["nombre_archivo", "titulo_principal", "diapositivas_json"]
                    }
                }
            },
            {
                "type": "function",
                "function": {
                    "name": "leer_presentacion_pptx",
                    "description": "Lee el texto y estructura de una presentación .pptx existente en el sandbox para resumirla o analizarla.",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "nombre_archivo": {
                                "type": "string",
                                "description": "Nombre del archivo .pptx en el sandbox."
                            }
                        },
                        "required": ["nombre_archivo"]
                    }
                }
            }
        ]

    def execute(self, tool_name: str, arguments: Dict[str, Any], user: str = "") -> str:
        if tool_name == "crear_presentacion_pptx":
            return self._crear_presentacion(arguments, user)
        elif tool_name == "leer_presentacion_pptx":
            return self._leer_presentacion(arguments, user)
        return f"Herramienta '{tool_name}' no soportada por PptxPlugin."

    def _crear_presentacion(self, args: Dict[str, Any], user: str) -> str:
        nombre = args.get("nombre_archivo", "presentacion.pptx").strip()
        if not nombre.lower().endswith(".pptx"):
            nombre += ".pptx"

        # Validación sandbox y path traversal
        os.makedirs(SANDBOX_DIR, exist_ok=True)
        ruta_salida = os.path.realpath(os.path.join(SANDBOX_DIR, nombre))
        if not ruta_salida.startswith(os.path.realpath(SANDBOX_DIR)):
            log_accion(user, "crear_presentacion_pptx", nombre, "BLOQUEADO: Path traversal")
            return "⛔ BLOQUEADO: Intento de Path Traversal."

        titulo_prin = args.get("titulo_principal", "Presentación")
        subtitulo = args.get("subtitulo", "Generado por Reaxy$ Agentic OS")
        diapositivas_raw = args.get("diapositivas_json", "[]")
        tema_nombre = args.get("tema", "azul_moderno")
        colores = TEMAS_COLOR.get(tema_nombre, TEMAS_COLOR["azul_moderno"])

        # Parsear diapositivas
        if isinstance(diapositivas_raw, str):
            try:
                diapositivas = json.loads(diapositivas_raw)
            except Exception:
                # Si viene texto plano o formateado, armamos una estructura de respaldo
                diapositivas = [{"titulo": "Puntos Clave", "puntos": [diapositivas_raw]}]
        else:
            diapositivas = diapositivas_raw

        try:
            prs = Presentation()
            # Dimensiones panorámicas 16:9 estándar
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)

            blank_layout = prs.slide_layouts[6]

            # 1. Diapositiva de Portada
            slide_portada = prs.slides.add_slide(blank_layout)
            tx_box = slide_portada.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.2))
            tf = tx_box.text_frame
            tf.word_wrap = True

            p1 = tf.paragraphs[0]
            p1.text = titulo_prin
            p1.font.bold = True
            p1.font.size = Pt(44)
            p1.font.color.rgb = colores["titulo"]
            p1.alignment = PP_ALIGN.LEFT

            p2 = tf.add_paragraph()
            p2.text = subtitulo
            p2.font.size = Pt(22)
            p2.font.color.rgb = colores["acento"]
            p2.space_before = Pt(14)
            p2.alignment = PP_ALIGN.LEFT

            # 2. Diapositivas de Contenido
            for d in diapositivas:
                slide = prs.slides.add_slide(blank_layout)
                
                # Título de diapositiva
                title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.2))
                tf_title = title_box.text_frame
                tf_title.word_wrap = True
                p_t = tf_title.paragraphs[0]
                p_t.text = d.get("titulo", "Tema")
                p_t.font.bold = True
                p_t.font.size = Pt(32)
                p_t.font.color.rgb = colores["titulo"]

                # Puntos / Contenido
                puntos = d.get("puntos", [])
                if isinstance(puntos, str):
                    puntos = [puntos]

                content_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.5))
                tf_content = content_box.text_frame
                tf_content.word_wrap = True

                for i, punto in enumerate(puntos):
                    p = tf_content.paragraphs[0] if i == 0 else tf_content.add_paragraph()
                    p.text = f"•  {punto}"
                    p.font.size = Pt(20)
                    p.font.color.rgb = colores["texto"]
                    p.space_before = Pt(12)

            prs.save(ruta_salida)
            log_accion(user, "crear_presentacion_pptx", nombre, f"OK ({len(diapositivas)} slides)")
            return f"✅ Presentación '{nombre}' creada exitosamente en {ruta_salida} con {len(diapositivas) + 1} diapositivas (Tema: {tema_nombre})."

        except Exception as e:
            log_accion(user, "crear_presentacion_pptx", nombre, f"ERROR: {e}")
            return f"❌ Error al crear presentación: {e}"

    def _leer_presentacion(self, args: Dict[str, Any], user: str) -> str:
        nombre = args.get("nombre_archivo", "")
        ruta_real = os.path.realpath(os.path.join(SANDBOX_DIR, nombre))
        if not ruta_real.startswith(os.path.realpath(SANDBOX_DIR)) or not os.path.exists(ruta_real):
            return f"❌ Archivo '{nombre}' no encontrado en el sandbox."

        try:
            prs = Presentation(ruta_real)
            resumen = [f"📊 Presentación: {nombre} ({len(prs.slides)} diapositivas)\n"]
            for idx, slide in enumerate(prs.slides):
                textos_slide = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        textos_slide.append(shape.text_frame.text.strip())
                resumen.append(f"--- Diapositiva {idx+1} ---")
                resumen.append("\n".join(textos_slide) if textos_slide else "[Diapositiva sin texto]")
            log_accion(user, "leer_presentacion_pptx", nombre, "OK")
            return "\n".join(resumen)[:3000]
        except Exception as e:
            log_accion(user, "leer_presentacion_pptx", nombre, f"ERROR: {e}")
            return f"❌ Error al leer la presentación: {e}"
